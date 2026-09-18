"""Writers that turn a :class:`DocumentSpec` into a real file on disk.

Real writers, real files: a PDF is paginated by PyMuPDF so page 14 genuinely
is page 14, a DOCX is a real OOXML package, a XLSX has real cells.  Extraction
downstream is therefore parsing, never replaying a fixture.
"""

from __future__ import annotations

import csv
import html
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Final

import pymupdf
from spectra_config.logging import get_logger

from .constants import (
    PDF_BODY_FONTSIZE,
    PDF_HEADING_FONTSIZE,
    PDF_LINE_HEIGHT,
    PDF_MARGIN,
    PDF_PAGE_HEIGHT,
    PDF_PAGE_WIDTH,
)
from .specs import (
    FORMAT_CSV,
    FORMAT_DOCX,
    FORMAT_HTML,
    FORMAT_MD,
    FORMAT_PDF,
    FORMAT_PPTX,
    FORMAT_TXT,
    FORMAT_XLSX,
    DocumentSpec,
)

log = get_logger(__name__)

#: Font sizes tried in order until the page body fits; failing all of them is an error.
PDF_FONT_LADDER: Final[tuple[float, ...]] = (PDF_BODY_FONTSIZE, 9.5, 8.5, 7.5, 6.5)
PDF_BODY_FONT: Final[str] = "helv"
PDF_HEADING_FONT: Final[str] = "hebo"
#: Fixed PDF timestamps keep regeneration reproducible.
PDF_FIXED_DATE: Final[str] = "D:20240101000000Z"
XLSX_SHEET_NAME: Final[str] = "transactions"
#: Fixed OOXML document properties and zip entry timestamps, so DOCX/PPTX/XLSX
#: packages regenerate byte-for-byte rather than embedding the wall clock.
OOXML_FIXED_MOMENT: Final[datetime] = datetime(2024, 1, 1, 0, 0, 0)
OOXML_ZIP_TIMESTAMP: Final[tuple[int, int, int, int, int, int]] = (1980, 1, 1, 0, 0, 0)
OOXML_ENTRY_ATTRIBUTES: Final[int] = 0o600 << 16


class DocumentWriteError(RuntimeError):
    """Raised when a document cannot be rendered to its target format."""


@dataclass(frozen=True)
class DocumentArtifact:
    """What the manifest records about one written document."""

    document_id: str
    slug: str
    path: str
    media_type: str
    title: str
    pages: int
    version: str
    version_status: str
    key_page: int | None
    key_finding: str | None
    tags: tuple[str, ...]
    entities: tuple[str, ...]
    occurred_at: str | None
    size_bytes: int
    section_headings: tuple[str, ...] = ()

    def page_of(self, heading_prefix: str) -> int | None:
        """1-based page whose heading starts with ``heading_prefix`` (PDFs only)."""
        if self.media_type != "application/pdf":
            return None
        lowered = heading_prefix.lower()
        for index, heading in enumerate(self.section_headings):
            if heading.lower().startswith(lowered):
                return index + 1
        return None


def _freeze_core_properties(properties: object) -> None:
    """Pin the OOXML core properties that would otherwise carry the wall clock."""
    for attribute, value in (
        ("created", OOXML_FIXED_MOMENT),
        ("modified", OOXML_FIXED_MOMENT),
        ("revision", 1),
        ("last_modified_by", "SPECTRA demo generator"),
    ):
        if hasattr(properties, attribute):
            setattr(properties, attribute, value)


def _normalise_ooxml(path: Path) -> None:
    """Rewrite an OOXML package with fixed entry timestamps and attributes."""
    with zipfile.ZipFile(path) as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, payload in entries:
            info = zipfile.ZipInfo(filename=name, date_time=OOXML_ZIP_TIMESTAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = OOXML_ENTRY_ATTRIBUTES
            target.writestr(info, payload)


def _banner(spec: DocumentSpec, document_id: str) -> str:
    return f"Document {document_id} | {spec.title} | version {spec.version} ({spec.version_status})"


def _write_pdf(spec: DocumentSpec, document_id: str, path: Path) -> int:
    document = pymupdf.open()
    total = len(spec.sections)
    try:
        for index, section in enumerate(spec.sections):
            page = document.new_page(width=PDF_PAGE_WIDTH, height=PDF_PAGE_HEIGHT)
            head_rect = pymupdf.Rect(PDF_MARGIN, PDF_MARGIN, PDF_PAGE_WIDTH - PDF_MARGIN, PDF_MARGIN + 60)
            page.insert_textbox(
                head_rect, section.heading, fontsize=PDF_HEADING_FONTSIZE, fontname=PDF_HEADING_FONT
            )
            body = "\n\n".join(section.paragraphs)
            footer = f"{_banner(spec, document_id)} | Page {index + 1} of {total}"
            text = f"{body}\n\n{footer}"
            body_rect = pymupdf.Rect(
                PDF_MARGIN,
                PDF_MARGIN + 70,
                PDF_PAGE_WIDTH - PDF_MARGIN,
                PDF_PAGE_HEIGHT - PDF_MARGIN,
            )
            if not _fit_text(page, body_rect, text):
                raise DocumentWriteError(f"page {index + 1} of {spec.slug} does not fit on one page")
        document.set_metadata(
            {
                "title": spec.title,
                "author": spec.author,
                "subject": ", ".join(spec.tags),
                "keywords": " ".join(spec.entities),
                "creationDate": PDF_FIXED_DATE,
                "modDate": PDF_FIXED_DATE,
                "producer": "SPECTRA demo generator",
                "creator": "SPECTRA demo generator",
            }
        )
        document.save(str(path), garbage=4, deflate=True, no_new_id=True, reproducible=True)
    finally:
        document.close()
    return total


def _fit_text(page: pymupdf.Page, rect: pymupdf.Rect, text: str) -> bool:
    for fontsize in PDF_FONT_LADDER:
        overflow = page.insert_textbox(
            rect, text, fontsize=fontsize, fontname=PDF_BODY_FONT, lineheight=PDF_LINE_HEIGHT / fontsize
        )
        if overflow >= 0:
            return True
    return False


def _write_docx(spec: DocumentSpec, document_id: str, path: Path) -> int:
    from docx import Document

    document = Document()
    document.add_heading(spec.title, level=1)
    document.add_paragraph(_banner(spec, document_id))
    for section in spec.sections:
        document.add_heading(section.heading, level=2)
        for paragraph in section.paragraphs:
            document.add_paragraph(paragraph)
    _freeze_core_properties(document.core_properties)
    document.save(str(path))
    _normalise_ooxml(path)
    return 1


def _write_pptx(spec: DocumentSpec, document_id: str, path: Path) -> int:
    from pptx import Presentation

    presentation = Presentation()
    title_layout = presentation.slide_layouts[0]
    bullet_layout = presentation.slide_layouts[1]
    opening = presentation.slides.add_slide(title_layout)
    opening.shapes.title.text = spec.title
    opening.placeholders[1].text = _banner(spec, document_id)
    for section in spec.sections:
        slide = presentation.slides.add_slide(bullet_layout)
        slide.shapes.title.text = section.heading
        frame = slide.placeholders[1].text_frame
        frame.text = section.paragraphs[0] if section.paragraphs else ""
        for paragraph in section.paragraphs[1:]:
            frame.add_paragraph().text = paragraph
    _freeze_core_properties(presentation.core_properties)
    presentation.save(str(path))
    _normalise_ooxml(path)
    return len(spec.sections) + 1


def _write_xlsx(spec: DocumentSpec, document_id: str, path: Path) -> int:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = XLSX_SHEET_NAME
    sheet.append([_banner(spec, document_id)])
    sheet.append(list(spec.headers))
    for row in spec.rows:
        sheet.append(list(row))
    _freeze_core_properties(workbook.properties)
    workbook.save(str(path))
    _normalise_ooxml(path)
    return 1


def _write_csv(spec: DocumentSpec, document_id: str, path: Path) -> int:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([f"# {_banner(spec, document_id)}"])
        writer.writerow(list(spec.headers))
        writer.writerows([list(row) for row in spec.rows])
    return 1


def _write_markdown(spec: DocumentSpec, document_id: str, path: Path) -> int:
    lines = [f"# {spec.title}", "", f"_{_banner(spec, document_id)}_", ""]
    for section in spec.sections:
        lines.extend([f"## {section.heading}", ""])
        lines.extend([f"{paragraph}\n" for paragraph in section.paragraphs])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return 1


def _write_html(spec: DocumentSpec, document_id: str, path: Path) -> int:
    parts = [
        "<!DOCTYPE html>",
        '<html lang="en"><head><meta charset="utf-8">',
        f"<title>{html.escape(spec.title)}</title></head><body>",
        f"<h1>{html.escape(spec.title)}</h1>",
        f"<p class=\"banner\">{html.escape(_banner(spec, document_id))}</p>",
    ]
    for section in spec.sections:
        parts.append(f"<h2>{html.escape(section.heading)}</h2>")
        parts.extend(f"<p>{html.escape(paragraph)}</p>" for paragraph in section.paragraphs)
    parts.append("</body></html>")
    path.write_text("\n".join(parts) + "\n", encoding="utf-8")
    return 1


def _write_txt(spec: DocumentSpec, document_id: str, path: Path) -> int:
    lines = [spec.title, _banner(spec, document_id), ""]
    for section in spec.sections:
        lines.extend([section.heading, "-" * len(section.heading)])
        lines.extend(section.paragraphs)
        lines.append("")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return 1


_WRITERS = {
    FORMAT_PDF: _write_pdf,
    FORMAT_DOCX: _write_docx,
    FORMAT_PPTX: _write_pptx,
    FORMAT_XLSX: _write_xlsx,
    FORMAT_CSV: _write_csv,
    FORMAT_MD: _write_markdown,
    FORMAT_HTML: _write_html,
    FORMAT_TXT: _write_txt,
}


def write_document(spec: DocumentSpec, document_id: str, directory: Path, relative_to: Path) -> DocumentArtifact:
    """Render one spec and return the manifest record for it."""
    writer = _WRITERS.get(spec.fmt)
    if writer is None:
        raise DocumentWriteError(f"no writer registered for format {spec.fmt!r}")
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / spec.filename
    try:
        pages = writer(spec, document_id, path)
    except DocumentWriteError:
        raise
    except Exception as exc:  # noqa: BLE001 - surfaced with the failing document
        raise DocumentWriteError(f"failed to write {spec.filename}: {exc}") from exc
    log.debug("demo_data.document_written", document_id=document_id, path=str(path), pages=pages)
    return DocumentArtifact(
        document_id=document_id,
        slug=spec.slug,
        path=path.relative_to(relative_to).as_posix(),
        media_type=spec.media_type,
        title=spec.title,
        pages=pages,
        version=spec.version,
        version_status=spec.version_status,
        key_page=spec.key_page,
        key_finding=spec.key_finding,
        tags=spec.tags,
        entities=spec.entities,
        occurred_at=spec.occurred_at.strftime("%Y-%m-%dT%H:%M:%SZ") if spec.occurred_at else None,
        size_bytes=path.stat().st_size,
        section_headings=tuple(section.heading for section in spec.sections),
    )

"""Office extraction: DOCX paragraphs/styles/tables, PPTX slides/notes, XLSX sheets."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import Any, Final

import openpyxl
from docx import Document as open_docx
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from pptx import Presentation as open_pptx
from spectra_config.logging import get_logger

from .base import ExtractedBlock, ExtractionError, markdown_table, normalise_text, sequence_blocks

log = get_logger(__name__)

XLSX_ROWS_PER_BLOCK: Final[int] = 50
XLSX_MAX_ROWS_PER_SHEET: Final[int] = 20_000
HEADING_STYLE_PREFIXES: Final[tuple[str, ...]] = ("heading", "title", "subtitle")


def extract_docx(path: str | Path) -> list[ExtractedBlock]:
    """Paragraphs (styles promoted to headings) and tables, in document order."""
    try:
        document = open_docx(str(path))
    except Exception as exc:
        raise ExtractionError(f"cannot open docx {Path(path).name}: {exc}") from exc

    blocks: list[ExtractedBlock] = []
    section: str | None = None
    for element in _iter_docx_body(document):
        if isinstance(element, DocxParagraph):
            block, section = _docx_paragraph(element, section)
            if block is not None:
                blocks.append(block)
            continue
        rendered = markdown_table([[cell.text for cell in row.cells] for row in element.rows])
        if rendered:
            blocks.append(ExtractedBlock(text=rendered, kind="table", section=section))
    return sequence_blocks(blocks)


def extract_pptx(path: str | Path) -> list[ExtractedBlock]:
    """One block per text frame plus the speaker notes; slide number is the page."""
    try:
        presentation = open_pptx(str(path))
    except Exception as exc:
        raise ExtractionError(f"cannot open pptx {Path(path).name}: {exc}") from exc

    blocks: list[ExtractedBlock] = []
    for index, slide in enumerate(presentation.slides):
        page = index + 1
        section = _slide_title(slide) or f"Slide {page}"
        blocks.append(ExtractedBlock(text=section, kind="heading", page=page, section=section))
        blocks.extend(_slide_body(slide, page, section))
        notes = _slide_notes(slide)
        if notes:
            blocks.append(
                ExtractedBlock(
                    text=notes, kind="caption", page=page, section=section, metadata={"origin": "notes"}
                )
            )
    return sequence_blocks(blocks)


def extract_xlsx(path: str | Path) -> list[ExtractedBlock]:
    """Each sheet becomes header-aware markdown tables, grouped by row windows."""
    try:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError(f"cannot open xlsx {Path(path).name}: {exc}") from exc

    blocks: list[ExtractedBlock] = []
    try:
        for index, sheet in enumerate(workbook.worksheets):
            blocks.extend(_sheet_blocks(sheet, index + 1))
    finally:
        workbook.close()
    return sequence_blocks(blocks)


def _iter_docx_body(document: Any) -> Iterator[DocxParagraph | DocxTable]:
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield DocxParagraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, document)


def _docx_paragraph(
    paragraph: DocxParagraph, section: str | None
) -> tuple[ExtractedBlock | None, str | None]:
    text = normalise_text(paragraph.text)
    if not text:
        return None, section
    style = (getattr(paragraph.style, "name", "") or "").strip().lower()
    if style.startswith(HEADING_STYLE_PREFIXES):
        return ExtractedBlock(text=text, kind="heading", section=text), text
    return ExtractedBlock(text=text, kind="paragraph", section=section), section


def _slide_title(slide: Any) -> str | None:
    try:
        title_shape = slide.shapes.title
    except Exception:  # pragma: no cover - layout without placeholders
        return None
    if title_shape is None or not getattr(title_shape, "has_text_frame", False):
        return None
    return normalise_text(title_shape.text_frame.text) or None


def _slide_body(slide: Any, page: int, section: str) -> list[ExtractedBlock]:
    blocks: list[ExtractedBlock] = []
    title = _slide_title(slide)
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = normalise_text(shape.text_frame.text)
        if not text or text == title:
            continue
        blocks.append(ExtractedBlock(text=text, kind="paragraph", page=page, section=section))
    return blocks


def _slide_notes(slide: Any) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        return normalise_text(slide.notes_slide.notes_text_frame.text)
    except Exception as exc:  # pragma: no cover - malformed notes part
        log.warning("pptx.notes_failed", error=str(exc))
        return ""


def _sheet_blocks(sheet: Any, page: int) -> list[ExtractedBlock]:
    title = str(getattr(sheet, "title", f"Sheet{page}"))
    try:
        rows = _sheet_rows(sheet)
    except Exception as exc:
        log.warning("xlsx.sheet_failed", sheet=title, error=str(exc))
        return []
    if not rows:
        return []

    header, body = rows[0], rows[1:]
    if not body:
        rendered = markdown_table([header])
        return [ExtractedBlock(text=rendered, kind="table", page=page, section=title)] if rendered else []

    blocks: list[ExtractedBlock] = []
    for start in range(0, len(body), XLSX_ROWS_PER_BLOCK):
        window = body[start : start + XLSX_ROWS_PER_BLOCK]
        rendered = markdown_table([header, *window])
        if not rendered:
            continue
        blocks.append(
            ExtractedBlock(
                text=rendered,
                kind="table",
                page=page,
                section=title,
                metadata={"sheet": title, "first_row": start + 2, "row_count": len(window)},
            )
        )
    return blocks


def _sheet_rows(sheet: Any) -> list[list[str]]:
    rows: list[list[str]] = []
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        if index >= XLSX_MAX_ROWS_PER_SHEET:
            log.warning("xlsx.sheet_truncated", sheet=str(sheet.title), limit=XLSX_MAX_ROWS_PER_SHEET)
            break
        values = ["" if value is None else str(value) for value in row]
        if any(value.strip() for value in values):
            rows.append(values)
    return rows
